import re
import openpyxl

# Read current lessons.ts to build translation mapping
translation_map = {}
with open("/Users/surachartlimrattanaphun/Desktop/My Project/EL/src/utils/lessons.ts", "r", encoding="utf-8") as f:
    ts_content = f.read()

# Extract all dialogue lines from lessons.ts
# Format: { speaker: '...', character: '...', text: '...', translation: '...', phonetic: '' }
matches = re.findall(r"text:\s*['\"](.*?)['\"].*?translation:\s*['\"](.*?)['\"]", ts_content)
for text, trans in matches:
    translation_map[text.strip()] = trans.strip()

print(f"Loaded {len(translation_map)} translations from lessons.ts")

wb = openpyxl.load_workbook("/Users/surachartlimrattanaphun/Desktop/My Project/EL/data/Word-Sentences-RolePlay_Unit1-5.xlsx")

# Let's inspect Unit 1 parsing
sheet1 = wb["Unit1 RolePlay"]
current_topic = ""
topics = {}
for r in range(3, sheet1.max_row + 1):
    val_a = sheet1.cell(r, 1).value
    val_b = sheet1.cell(r, 2).value
    
    if val_a and not val_b:
        current_topic = val_a.strip()
        topics[current_topic] = []
    elif val_b and current_topic:
        # Dialogue line(s)
        topics[current_topic].append((val_a, val_b))

print("\n--- Unit 1 Parsed Topics ---")
for t, lines in topics.items():
    print(f"Topic: {t} ({len(lines)} rows)")
    for row_num, (lbl, text) in enumerate(lines[:2]):
        print(f"  Row {row_num}: [{lbl}] -> {repr(text)}")
