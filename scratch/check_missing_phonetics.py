import re
import openpyxl

# Read lessons.ts
with open("/Users/surachartlimrattanaphun/Desktop/My Project/EL/src/utils/lessons.ts", "r", encoding="utf-8") as f:
    ts_content = f.read()

# Build translation and phonetic lookup map
trans_map = {}
phone_map = {}

vocab_matches = re.findall(r"word:\s*['\"](.*?)['\"].*?phonetic:\s*['\"](.*?)['\"].*?translation:\s*['\"](.*?)['\"]", ts_content)
for w, p, t in vocab_matches:
    trans_map[w.strip().lower()] = t.strip()
    phone_map[w.strip().lower()] = p.strip()

sentence_matches = re.findall(r"text:\s*['\"](.*?)['\"].*?translation:\s*['\"](.*?)['\"]", ts_content)
for s, t in sentence_matches:
    trans_map[s.strip().lower()] = t.strip()

dialogue_matches = re.findall(r"text:\s*['\"](.*?)['\"].*?translation:\s*['\"](.*?)['\"].*?phonetic:\s*['\"](.*?)['\"]", ts_content)
for text, trans, phone in dialogue_matches:
    text_clean = text.strip().lower()
    if trans.strip():
        trans_map[text_clean] = trans.strip()
    if phone.strip():
        phone_map[text_clean] = phone.strip()

# Custom translations dictionary including the missing translations we identified
custom_translations = {
    "what's this?": "นี่คืออะไรเหรอ?",
    "what's that?": "นั่นคืออะไรเหรอ?",
    "what is this?": "นี่คืออะไรเหรอ?",
    "yes, it is.": "ใช่แล้วจ้า",
    "no, it isn't.": "ไม่ใช่จ้า",
    "no, it is not.": "ไม่ใช่จ้า",
    "yes, I do.": "ใช่แล้วจ้า / ชอบครับ/ค่ะ",
    "no, I don't.": "ไม่ใช่จ้า / ไม่ชอบครับ/ค่ะ",
    "thank you!": "ขอบคุณนะ!",
    "thanks a lot.": "ขอบคุณมาก ๆ เลย",
    "good morning!": "อรุณสวัสดิ์ครับ/ค่ะ!",
    "where do you live, b?": "เธออาศัยอยู่ที่ไหนเหรอ บี?",
    "i live in trang.": "ฉันอาศัยอยู่ในจังหวัดตรังจ้า",
    "are you thai?": "เธอเป็นคนไทยหรือเปล่า?",
    "yes, i am thai. and you?": "ใช่แล้ว ฉันเป็นคนไทย แล้วเธอล่ะ?",
    "my family is chinese.": "ครอบครัวของฉันเป็นคนจีนจ้า",
    "is your friend muslim?": "เพื่อนของเธอเป็นชาวมุสลิมใช่ไหม?",
    "yes, she is. she loves her family.": "ใช่แล้วจ้า เธอรักครอบครัวของเธอมากเลย",
    "what is that purple flower?": "ดอกไม้สีม่วงนั่นคือดอกอะไรเหรอ?",
    "it is sri trang flower. it is beautiful!": "มันคือดอกศรีตรังจ้า สวยงามมากเลย!",
    "look! many rubber trees.": "ดูนั่นสิ! ต้นยางพาราเยอะแยะเลย",
    "wow! my grandfather has a rubber farm.": "ว้าว! คุณปู่ของฉันมีสวนยางพาราด้วยแหละ",
    "who is phraya ratsada?": "พระยารัษฎาฯ คือใครเหรอ?",
    "he was a great governor of trang.": "ท่านเคยเป็นผู้ว่าราชการเมืองตรังที่ยิ่งใหญ่จ้า",
    "do you like trang roasted pork?": "เธอชอบหมูย่างเมืองตรังไหม?",
    "yes, i do. it is very yummy!": "ชอบมากเลยจ้า อร่อยมาก ๆ!",
    "let's eat dim sum this morning.": "เช้านี้มากินติ่มซำกันเถอะ",
    "great idea! i love dim sum.": "เป็นความคิดที่ดีเลย! ฉันชอบติ่มซำมาก",
    "this trang cake is for you.": "เค้กเมืองตรังชิ้นนี้สำหรับเธอนะ",
    "thank you! i like sweet cake.": "ขอบคุณนะ! ฉันชอบกินเค้กหวาน ๆ จ้า",
    "where is the clock tower?": "หอนาฬิกาตั้งอยู่ที่ไหนเหรอ?",
    "it is in trang city. it is very big!": "อยู่ในเมืองตรังจ้า มันใหญ่โตมากเลย!",
    "do you know koh kradan?": "เธอรู้จักเกาะกระดานไหม?",
    "yes! the sea is so beautiful there.": "รู้จักสิ! ทะเลที่นั่นสวยงามมากเลยล่ะ",
    "let's go to emerald cave.": "ไปเที่ยวถ้ำมรกตกันเถอะ",
    "is it scary?": "มันน่ากลัวไหมอะ?",
    "no, it is a beautiful cave!": "ไม่หรอก มันเป็นถ้ำที่สวยงามมากเลย!",
    "i like ton te waterfall.": "ฉันชอบน้ำตกโตนเตะจ้า",
    "me too! the water is very cold.": "ฉันก็ชอบเหมือนกัน! น้ำเย็นเจี๊ยบเลย",
    "what is the underwater wedding?": "งานวิวาห์ใต้สมุทรคืออะไรเหรอ?",
    "it is a famous festival in trang. people marry under the sea!": "มันคือเทศกาลที่มีชื่อเสียงของเมืองตรังจ้า คนเขาแต่งงานกันใต้ท้องทะเลน่ะ!",
    "what is that in the wind?": "นั่นอะไรอยู่ในสายลมเหรอ?",
    "it is a looklom windmill. it goes spin, spin, spin!": "มันคือกังหันลมลูกลมจ้า มันหมุน ติ้ว ติ้ว ติ้ว เลย!",
    
    # Missing translations identified
    "can you spell it? (milk)": "เธอสะกดคำว่านม (Milk) ได้ไหม?",
    "yes, i did / of course": "ใช่แล้วครับ/ค่ะ / แน่นอนอยู่แล้ว!",
    "my parents have arrived!": "ผู้ปกครองของฉันมาถึงแล้วครับ/ค่ะ!",
    "yes, it is. it's beautiful!": "ใช่แล้วล่ะ มันสวยงามมากเลย!",
}

# Custom phonetics dictionary to provide accurate pronunciation guides for kids!
custom_phonetics = {
    # Hello / greetings
    "good morning!": "กูด มอร์นิง!",
    "what's your name?": "ว็อทส์ ยัวร์ เนม?",
    "my name is thanaphob": "มาย เนม อีส ธนภพ",
    "what's your nickname?": "ว็อทส์ ยัวร์ นิกเนม?",
    "my nickname is phob": "มาย นิกเนม อีส ภพ",
    "how old are you?": "ฮาว โอลด์ อาร์ ยู?",
    "i'm eight years old.": "ไอม์ เอท เยียร์ส โอลด์",
    "how are you?": "ฮาว อาร์ ยู?",
    "i am very well.": "ไอ แอม เวนี่ เวลล์",
    "how did you come to school?": "ฮาว ดิด ยู คัม ทู สคูล?",
    "i came to school by car.": "ไอ เคม ทู สคูล บาย คาร์",
    "did you have breakfast?": "ดิด ยู แฮฟ เบรคฟาสต์?",
    "yes, i did.": "เยส ไอ ดิด",
    "what grade are you in?": "ว็อท เกรด อาร์ ยู อิน?",
    "i'm grade in primary 2": "ไอม์ เกรด อิน ไพรมารี ทู",
    "what day is today?": "ว็อท เดย์ อีส ทูเดย์?",
    "today is monday.": "ทูเดย์ อีส มันเดย์",
    "how do you come to school?": "ฮาว ดู ยู คัม ทู สคูล?",
    "i come to school by car.": "ไอ คัม ทู สคูล บาย คาร์",
    "how many people are there in your room today?": "ฮาว เมนี่ พีเพิล อาร์ แดร์ อิน ยัวร์ รูม ทูเดย์?",
    "there are 31 people in my room.": "แดร์ อาร์ เทอร์ตี้วัน พีเพิล อิน มาย รูม",
    "are you ready to dance?": "อาร์ ยู เรดี้ ทู แดนซ์?",
    "yes ,we are.": "เยส วี อาร์",
    "yes, we are.": "เยส วี อาร์",
    "o.k. that's all for now. walk to the class.": "โอเค แดทส์ ออล ฟอร์ นาว วอล์ค ทู เดอะ คลาส",
    "thank you teacher.": "แธงค์ ยู ทีเชอร์",
    "thank you, teacher": "แธงค์ ยู ทีเชอร์",
    
    # Milk topic
    "what is this?": "ว็อท อีส ดิส?",
    "this is milk.": "ดิส อีส มิลค์",
    "can you spell it?": "แคน ยู สเปล อิท?",
    "m-i-l-k": "เอ็ม ไอ แอล เค",
    "please, cut the bag for me.": "พลีส คัท เดอะ แบ็ก ฟอร์ มี",
    "yes, of course": "เยส ออฟ คอร์ส",
    "what color is milk?": "ว็อท คัลเลอร์ อีส มิลค์?",
    "it's white.": "อิทส์ ไวท์",
    "do you like milk?": "ดู ยู ไลค์ มิลค์?",
    "yes, i do.": "เยส ไอ ดู",
    "can i have a straw please?": "แคน ไอ แฮฟ อะ สตรอว์ พลีส?",
    "yes, you can.": "เยส ยู แคน",
    
    # Unit 2 Classroom
    "what’s this?": "ว็อทส์ ดิส?",
    "i don’t know. is it a pencil?": "ไอ โดนท์ โนว์ อีส อิท อะ เพนซิล?",
    "no, it isn’t. it’s a ruler.": "โน อิท อิสเซินท์ อิทส์ อะ รูเลอร์",
    "is this a school bag?": "อีส ดิส อะ สคูล แบ็ก?",
    "yes, it is. it's beautiful!": "เยส อิท อีส อิทส์ บิวตี้ฟูล",
    "what’s that on the wall?": "ว็อทส์ แดท ออน เดอะ วอลล์?",
    "it’s a clock.": "อิทส์ อะ คล็อก",
    "oh, it’s 9 o’clock.": "โอ้ อิทส์ ไนน์ โอ คล็อก",
    "do you have an eraser?": "ดู ยู แฮฟ แอน อีเรเซอร์?",
    "yes, i do. here you are.": "เยส ไอ ดู เฮียร์ ยู อาร์",
    "thank you!": "แธงค์ ยู!",
    "can i borrow your crayon, please?": "แคน ไอ บอร์โรว์ ยัวร์ เครยอน พลีส?",
    "sure! i have a red crayon.": "ชัวร์! ไอ แฮฟ อะ เรด เครยอน",
    "thanks a lot.": "แธงค์ อะ ล็อท",
    "where is my notebook?": "แวร์ อีส มาย โน้ตบุ๊ก?",
    "look! it’s under your chair.": "ลุค! อิทส์ อัญเดอร์ ยัวร์ แชร์",
    "oh, thank you.": "โอ้ แธงค์ ยู",
    "is my pen on the table?": "อีส มาย เพน ออน เดอะ เทเบิล?",
    "no, it isn’t. it’s in your pencil case.": "โน อิท อิสเซินท์ อิทส์ อิน ยัวร์ เพนซิล เคส",
    "please close the window. it’s raining!": "พลีส โคลส เดอะ วินโดว์ อิทส์ เรนนิง!",
    "ok, i can do it.": "โอเค ไอ แคน ดู อิท",
    "open the door, please.": "โอเพน เดอะ ดอร์ พลีส",
    "time for lunch! let’s go.": "ไทม์ ฟอร์ ลันช์ เล็ทส์ โก",
    "great! i am hungry.": "เกรท! ไอ แอม ฮังกรี",
    "what do you have for lunch?": "ว็อท ดู ยู แฮฟ ฟอร์ ลันช์?",
    "i have rice and chicken.": "ไอ แฮฟ ไรซ์ แอนด์ ชิคเก้น",
    "let’s go to the playground.": "เล็ทส์ โก ทู เดอะ เพลย์กราวด์",
    "ok! let’s play a game.": "โอเค! เล็ทส์ เพลย์ อะ เกม",
    "look at that big ball!": "ลุค แอท แดท บิ๊ก บอลล์!",
    "is it your ball?": "อีส อิท ยัวร์ บอลล์?",
    "no, it isn’t. it’s sam’s ball.": "โน อิท อิสเซินท์ อิทส์ แซมส์ บอลล์",
    "shh... be quiet in the library.": "ชู่ บี ไควเอท อิน เดอะ ไลบรารี",
    "sorry. what is that book?": "ซอรี่ ว็อท อีส แดท บุ๊ก?",
    "it’s a book about animals.": "อิทส์ อะ บุ๊ก อะเบาท์ แอนิมอลส์",
    "do we have homework today?": "ดู วี แฮฟ โฮมเวิร์ก ทูเดย์?",
    "yes. we have english homework.": "เยส วี แฮฟ อิงลิช โฮมเวิร์ก",
    "let’s do it together!": "เล็ทส์ ดู อิท ทูเกเธอร์!",
}

def get_phonetic_guide(text):
    t_clean = text.strip().lower().replace("’", "'").replace('"', "'").strip()
    # Normalize ending punctuation
    t_clean_nopunct = re.sub(r'[.!?]$', '', t_clean).strip()
    
    if t_clean in custom_phonetics:
        return custom_phonetics[t_clean]
    if t_clean_nopunct in custom_phonetics:
        return custom_phonetics[t_clean_nopunct]
        
    # Check phone_map from lessons.ts
    if t_clean in phone_map and phone_map[t_clean]:
        return phone_map[t_clean]
    if t_clean_nopunct in phone_map and phone_map[t_clean_nopunct]:
        return phone_map[t_clean_nopunct]
        
    # Generate automatic simple phonetics or keep empty
    # For now, let's return a placeholder or empty so we know what is missing
    return ""

wb = openpyxl.load_workbook("/Users/surachartlimrattanaphun/Desktop/My Project/EL/data/Word-Sentences-RolePlay_Unit1-5.xlsx")

missing_phonetics_count = 0

for unit_id in range(1, 6):
    sheetname = f"Unit{unit_id} RolePlay"
    sheet = wb[sheetname]
    
    row_idx = 3
    while row_idx <= sheet.max_row:
        val_b = sheet.cell(row_idx, 2).value
        if not val_b:
            row_idx += 1
            continue
            
        # Parse dialogue turns
        dialogue_text = val_b.strip()
        speaker_pattern = re.compile(r'([ABTS]):\s*(.*?)(?=\s+[ABTS]:|$)')
        matches = speaker_pattern.findall(dialogue_text)
        
        for spk, txt in matches:
            txt = txt.strip()
            phone = get_phonetic_guide(txt)
            if not phone:
                missing_phonetics_count += 1
                print(f"Missing Phonetic [Unit {unit_id} Row {row_idx}]: {repr(txt)}")
                
        row_idx += 1

print(f"\nTotal Missing Phonetics: {missing_phonetics_count}")
