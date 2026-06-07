import re
import openpyxl

# Read lessons.ts
with open("/Users/surachartlimrattanaphun/Desktop/My Project/EL/src/utils/lessons.ts", "r", encoding="utf-8") as f:
    ts_content = f.read()

# Build translation and phonetic lookup map
trans_map = {}
phone_map = {}

# Parse vocab items translation/phonetic
# e.g., { word: 'Nickname', phonetic: 'นิกเนม', translation: 'ชื่อเล่น', emoji: '😊' }
vocab_matches = re.findall(r"word:\s*['\"](.*?)['\"].*?phonetic:\s*['\"](.*?)['\"].*?translation:\s*['\"](.*?)['\"]", ts_content)
for w, p, t in vocab_matches:
    trans_map[w.strip().lower()] = t.strip()
    phone_map[w.strip().lower()] = p.strip()

# Parse sentence items
sentence_matches = re.findall(r"text:\s*['\"](.*?)['\"].*?translation:\s*['\"](.*?)['\"]", ts_content)
for s, t in sentence_matches:
    trans_map[s.strip().lower()] = t.strip()

# Parse dialogue matches
dialogue_matches = re.findall(r"text:\s*['\"](.*?)['\"].*?translation:\s*['\"](.*?)['\"].*?phonetic:\s*['\"](.*?)['\"]", ts_content)
for text, trans, phone in dialogue_matches:
    text_clean = text.strip().lower()
    if trans.strip():
        trans_map[text_clean] = trans.strip()
    if phone.strip():
        phone_map[text_clean] = phone.strip()

# Let's add manual custom fixes if lookups fail or have minor punctuation differences
custom_translations = {
    "what's this?": "นี่คืออะไรเหรอ?",
    "what's that?": "นั่นคืออะไรเหรอ?",
    "what is this?": "นี่คืออะไรเหรอ?",
    "yes, it is.": "ใช่แล้วจ้า",
    "no, it isn't.": "ไม่ใช่จ้า",
    "no, it is not.": "ไม่ใช่จ้า",
    "yes, I do.": "ใช่แล้วจ้า / ชอบครับ/ค่ะ",
    "no, I don't.": "ไม่ใช่จ้า / ไม่ชอบครับ/ค่ะ",
    "thank you!": "ขอบคุณนะ!",
    "thanks a lot.": "ขอบคุณมาก ๆ เลย",
    "good morning!": "อรุณสวัสดิ์ครับ/ค่ะ!",
    "where do you live, b?": "เธออาศัยอยู่ที่ไหนเหรอ บี?",
    "i live in trang.": "ฉันอาศัยอยู่ในจังหวัดตรังจ้า",
    "are you thai?": "เธอเป็นคนไทยหรือเปล่า?",
    "yes, i am thai. and you?": "ใช่แล้ว ฉันเป็นคนไทย แล้วเธอล่ะ?",
    "my family is chinese.": "ครอบครัวของฉันเป็นคนจีนจ้า",
    "is your friend muslim?": "เพื่อนของเธอเป็นชาวมุสลิมใช่ไหม?",
    "yes, she is. she loves her family.": "ใช่แล้วจ้า เธอรักครอบครัวของเธอมากเลย",
    "what is that purple flower?": "ดอกไม้สีม่วงนั่นคือดอกอะไรเหรอ?",
    "it is sri trang flower. it is beautiful!": "มันคือดอกศรีตรังจ้า สวยงามมากเลย!",
    "look! many rubber trees.": "ดูนั่นสิ! ต้นยางพาราเยอะแยะเลย",
    "wow! my grandfather has a rubber farm.": "ว้าว! คุณปู่ของฉันมีสวนยางพาราด้วยแหละ",
    "who is phraya ratsada?": "พระยารัษฎาฯ คือใครเหรอ?",
    "he was a great governor of trang.": "ท่านเคยเป็นผู้ว่าราชการเมืองตรังที่ยิ่งใหญ่จ้า",
    "do you like trang roasted pork?": "เธอชอบหมูย่างเมืองตรังไหม?",
    "yes, i do. it is very yummy!": "ชอบมากเลยจ้า อร่อยมาก ๆ!",
    "let's eat dim sum this morning.": "เช้านี้มากินติ่มซำกันเถอะ",
    "great idea! i love dim sum.": "เป็นความคิดที่ดีเลย! ฉันชอบติ่มซำมาก",
    "this trang cake is for you.": "เค้กเมืองตรังชิ้นนี้สำหรับเธอนะ",
    "thank you! i like sweet cake.": "ขอบคุณนะ! ฉันชอบกินเค้กหวาน ๆ จ้า",
    "where is the clock tower?": "หอนาฬิกาตั้งอยู่ที่ไหนเหรอ?",
    "it is in trang city. it is very big!": "อยู่ในเมืองตรังจ้า มันใหญ่โตมากเลย!",
    "do you know koh kradan?": "เธอรู้จักเกาะกระดานไหม?",
    "yes! the sea is so beautiful there.": "รู้จักสิ! ทะเลที่นั่นสวยงามมากเลยล่ะ",
    "let's go to emerald cave.": "ไปเที่ยวถ้ำมรกตกันเถอะ",
    "is it scary?": "มันน่ากลัวไหมอะ?",
    "no, it is a beautiful cave!": "ไม่หรอก มันเป็นถ้ำที่สวยงามมากเลย!",
    "i like ton te waterfall.": "ฉันชอบน้ำตกโตนเตะจ้า",
    "me too! the water is very cold.": "ฉันก็ชอบเหมือนกัน! น้ำเย็นเจี๊ยบเลย",
    "what is the underwater wedding?": "งานวิวาห์ใต้สมุทรคืออะไรเหรอ?",
    "it is a famous festival in trang. people marry under the sea!": "มันคือเทศกาลที่มีชื่อเสียงของเมืองตรังจ้า คนเขาแต่งงานกันใต้ท้องทะเลน่ะ!",
    "what is that in the wind?": "นั่นอะไรอยู่ในสายลมเหรอ?",
    "it is a looklom windmill. it goes spin, spin, spin!": "มันคือกังหันลมลูกลมจ้า มันหมุน ติ้ว ติ้ว ติ้ว เลย!",
}

def lookup_translation(text):
    t_clean = text.strip().lower()
    # Normalize punctuation and quotes
    t_norm = t_clean.replace("’", "'").replace("`", "'").replace('"', "'").strip()
    if t_norm in custom_translations:
        return custom_translations[t_norm]
    if t_clean in trans_map:
        return trans_map[t_clean]
    if t_norm in trans_map:
        return trans_map[t_norm]
    # Remove final punctuation
    t_nopunct = re.sub(r'[.!?]$', '', t_norm).strip()
    if t_nopunct in trans_map:
        return trans_map[t_nopunct]
    # Fallback to direct lookup
    for k, v in trans_map.items():
        if k.replace("’", "'").replace('"', "'").strip() == t_norm:
            return v
    return ""

def lookup_phonetic(text):
    t_clean = text.strip().lower()
    t_norm = t_clean.replace("’", "'").replace('"', "'").strip()
    if t_clean in phone_map:
        return phone_map[t_clean]
    if t_norm in phone_map:
        return phone_map[t_norm]
    return ""

wb = openpyxl.load_workbook("/Users/surachartlimrattanaphun/Desktop/My Project/EL/data/Word-Sentences-RolePlay_Unit1-5.xlsx")

parsed_lessons_topics = {}

for unit_id in range(1, 6):
    sheetname = f"Unit{unit_id} RolePlay"
    sheet = wb[sheetname]
    
    topics_list = []
    current_topic_name = "General"
    current_dialogue = []
    
    row_idx = 3
    while row_idx <= sheet.max_row:
        val_a = sheet.cell(row_idx, 1).value
        val_b = sheet.cell(row_idx, 2).value
        
        if not val_a and not val_b:
            row_idx += 1
            continue
            
        if val_a and not val_b:
            # New topic header (Unit 1, 2, 4, 5)
            if current_dialogue:
                topics_list.append({
                    'title': current_topic_name,
                    'dialogue': current_dialogue
                })
                current_dialogue = []
            current_topic_name = val_a.strip()
            row_idx += 1
            continue
            
        # Dialogue row
        lbl = val_a.strip() if val_a else ""
        dialogue_text = val_b.strip() if val_b else ""
        
        if unit_id == 3:
            # For Unit 3, each row itself is a topic
            if current_dialogue:
                topics_list.append({
                    'title': current_topic_name,
                    'dialogue': current_dialogue
                })
            current_topic_name = lbl
            current_dialogue = []
            
        # Parse dialogue turns
        turns = []
        speaker_pattern = re.compile(r'([ABTS]):\s*(.*?)(?=\s+[ABTS]:|$)')
        matches = speaker_pattern.findall(dialogue_text)
        
        for spk, txt in matches:
            spk = spk.strip()
            txt = txt.strip()
            speaker_role = 'A' if spk in ['A', 'T'] else 'B'
            char_role = 'dino' if speaker_role == 'A' else 'bear'
            
            trans = lookup_translation(txt)
            phone = lookup_phonetic(txt)
            
            # Print warnings for missing translations
            if not trans:
                print(f"WARNING: Unit {unit_id} Missing translation for: {repr(txt)}")
                
            turns.append({
                'speaker': speaker_role,
                'character': char_role,
                'text': txt,
                'translation': trans,
                'phonetic': phone
            })
            
        current_dialogue.extend(turns)
        row_idx += 1
        
    # Append the last topic
    if current_dialogue:
        topics_list.append({
            'title': current_topic_name,
            'dialogue': current_dialogue
        })
        
    parsed_lessons_topics[unit_id] = topics_list
    print(f"Unit {unit_id}: parsed {len(topics_list)} topics")
