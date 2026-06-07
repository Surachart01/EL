import openpyxl

wb = openpyxl.load_workbook("/Users/surachartlimrattanaphun/Desktop/My Project/EL/data/Word-Sentences-RolePlay_Unit1-5.xlsx")

for unit_id in range(1, 6):
    sheetname = f"Unit{unit_id} RolePlay"
    if sheetname not in wb.sheetnames:
         print(f"Sheet {sheetname} not found!")
         continue
    sheet = wb[sheetname]
    print(f"\n==================== {sheetname} ====================")
    for row_idx in range(1, sheet.max_row + 1):
        vals = [cell.value for cell in sheet[row_idx]]
        if not any(vals):
            continue
        print(f"Row {row_idx}: {vals[:4]}")
