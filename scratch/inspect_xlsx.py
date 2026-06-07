import openpyxl

wb = openpyxl.load_workbook("/Users/surachartlimrattanaphun/Desktop/My Project/EL/data/Word-Sentences-RolePlay_Unit1-5.xlsx")
print("Sheets:", wb.sheetnames)

for sheetname in wb.sheetnames:
    sheet = wb[sheetname]
    print(f"\nSheet: {sheetname}")
    for i in range(1, 10):
        row = [cell.value for cell in sheet[i]]
        if any(row):
            print(f"Row {i}:", row[:6])
