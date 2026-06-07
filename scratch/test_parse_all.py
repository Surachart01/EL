import re
import openpyxl

wb = openpyxl.load_workbook("/Users/surachartlimrattanaphun/Desktop/My Project/EL/data/Word-Sentences-RolePlay_Unit1-5.xlsx")

for unit_id in range(1, 6):
    sheetname = f"Unit{unit_id} RolePlay"
    sheet = wb[sheetname]
    print(f"\n==================== {sheetname} ====================")
    
    current_topic = "General"
    row_idx = 3
    while row_idx <= sheet.max_row:
        val_a = sheet.cell(row_idx, 1).value
        val_b = sheet.cell(row_idx, 2).value
        
        if not val_a and not val_b:
            row_idx += 1
            continue
            
        if val_a and not val_b:
            current_topic = val_a.strip()
            print(f"Topic Header: {current_topic}")
            row_idx += 1
            continue
            
        # Dialogue row
        lbl = val_a.strip() if val_a else ""
        dialogue_text = val_b.strip() if val_b else ""
        
        # If Unit 3, since there are no header rows, each row acts as a topic!
        if unit_id == 3:
            current_topic = lbl
            print(f"Unit 3 Topic: {current_topic}")
            
        # Parse dialogue turns
        turns = []
        # Support T/S for Unit 1, and A/B for others
        speaker_pattern = re.compile(r'([ABTS]):\s*(.*?)(?=\s+[ABTS]:|$)')
        
        # Find all turns in this cell
        matches = speaker_pattern.findall(dialogue_text)
        for spk, txt in matches:
            spk = spk.strip()
            txt = txt.strip()
            # Map T -> A, S -> B
            speaker_role = 'A' if spk in ['A', 'T'] else 'B'
            char_role = 'dino' if speaker_role == 'A' else 'bear'
            turns.append({
                'speaker': speaker_role,
                'character': char_role,
                'text': txt
            })
            
        print(f"  Row {row_idx} [{lbl}] -> {len(turns)} turns")
        for t in turns[:2]:
            print(f"    {t['speaker']} ({t['character']}): {t['text']}")
            
        row_idx += 1
